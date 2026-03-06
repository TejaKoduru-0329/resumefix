import re
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
import pdfplumber
from docx import Document
from django.conf import settings
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, ListFlowable, Spacer, KeepTogether, ListItem
)
from reportlab.platypus.flowables import HRFlowable
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.lib import colors
from .template_config import RESUME_TEMPLATES


# 1. TEXT EXTRACTION
def extract_text_from_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def extract_text_from_docx(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


# 2. AI BRAIN
def get_ai_optimized_resume(resume_text, job_description):
    url = "https://api.groq.com/openai/v1/chat/completions"

    prompt = f"""You are an expert ATS resume rewriter. Rewrite the resume to maximize ATS score by matching the Job Description. Output ONLY the resume content, nothing else.

STRICT RULES:
- Plain text only. No HTML, markdown, or emojis.
- Line 1: Candidate NAME only
- Line 2: All contact details separated by " | " on ONE single line. Include ONLY details that are explicitly present in the original resume. Do NOT add "Not Available", "N/A", or any placeholder for missing fields. If DOB is not in resume, skip it entirely.
- Use **HEADING** format ONLY for section headings, NOT for job titles or project names
- Use • for ALL bullets. Never use - or numbers as bullets.
- No blank lines between sections

ATS OPTIMIZATION (CRITICAL):
- Carefully extract ALL important keywords, skills, tools, technologies, and action verbs from the Job Description
- Naturally inject these JD keywords into: Summary/Objective, Work Experience bullets, Project bullets, Technical Skills
- Rewrite ALL bullet points using strong action verbs from the JD (e.g. Developed, Implemented, Optimized, Collaborated)
- Summary/Objective MUST be tailored to the exact job role and requirements mentioned in the JD
- Do NOT fabricate experience — only enhance and rephrase existing content using JD language
- Add JD-relevant skills to Technical Skills if they are reasonably implied by the candidate's background
- Use EXACT keyword spelling from JD (e.g. if JD says "JavaScript" use "JavaScript" not "JS") 
- Repeat important JD keywords multiple times across different sections                           

SECTION ORDER (include only if applicable):
**CAREER OBJECTIVE** (if fresher or intern with no full-time experience) — Start with "To obtain..." or "Seeking a position...", focus on learning, academic background, and eagerness to contribute. Tailor to JD job role. ONE paragraph, no bullets, 3-4 lines.
**PROFESSIONAL SUMMARY** (ONLY if original resume has full-time work experience with company names and job titles) — Calculate exact years/months from the dates in the resume. Start with "X years of experience..." using ONLY the actual duration from resume. Do NOT fabricate or assume experience. ONE paragraph, no bullets, 3-4 lines.

IMPORTANT: If the original resume has NO full-time work experience (only internships or nothing), use **CAREER OBJECTIVE** instead. NEVER use PROFESSIONAL SUMMARY for freshers or interns.

**WORK EXPERIENCE** — include ONLY if the original resume explicitly contains internships, traineeships, apprenticeships, or full-time job roles with company names. If the original resume has NO work experience at all, DO NOT include this section. NEVER write "No relevant experience" or any placeholder text.
**EDUCATION** — STRICT 2 LINE FORMAT PER ENTRY:
Line 1: Institution Name, Location Year (ALL on ONE line, no line breaks)
Line 2: Degree | CGPA: value (OR GPA: value OR Percentage: value — use exactly what resume has, never combine)
NEVER put GPA/CGPA/Percentage on a separate third line.
NEVER split institution and year onto different lines.
**TECHNICAL SKILLS** — add JD-relevant skills implied by candidate background. Format: "Category: value1, value2"
**PROJECTS** — project title on one line (keep original title). Rewrite every bullet point to:
- Replace generic descriptions with JD-specific tools, technologies, and methodologies
- If project used any software/tools, mention them with JD-relevant context
- Connect project outcomes directly to JD requirements using exact JD keywords
- Every bullet MUST contain at least 2 JD keywords
**SOFT SKILLS** — COMPLETELY REWRITE using ONLY soft skills explicitly mentioned or strongly implied in the Job Description. Do NOT copy soft skills from the original resume. Extract communication, teamwork, leadership, analytical, and other soft skills from JD language and list them as • bullets.
**LANGUAGES KNOWN** — • bullets, only languages from original resume
**CERTIFICATIONS** — only if present in original resume, omit entirely if none

Resume:
{resume_text}

Job Description:
{job_description}
"""
    headers = {
        "Authorization": f"Bearer {settings.GROQ_API_KEY}",
        "Content-Type": "application/json"
    }

    session = requests.Session()
    retry = Retry(
        total=3,
        backoff_factor=2,
        status_forcelist=[429, 500, 502, 503, 504]
    )

    adapter = HTTPAdapter(max_retries=retry)
    session.mount("https://", adapter)
    
    payload = {
        "model": "llama-3.3-70b-versatile",
        "messages": [
            {
                "role": "system",
                "content": "You are an ATS optimization expert. Your ONLY job is to maximize keyword matching between resume and job description. Inject as many exact JD keywords as possible into every section."
            },

            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1,
        "max_tokens": 2048
    }
    
    response = requests.post(url, json=payload, headers=headers, timeout=(15, 120))
    data = response.json()

    if "choices" not in data:
        raise Exception(data.get("error", "Groq API error"))

    return data["choices"][0]["message"]["content"]


# 3. PDF GENERATION
def generate_ats_pdf(text, output_path, template_key="classic"):

    style_cfg = RESUME_TEMPLATES.get(template_key, RESUME_TEMPLATES["classic"])

    # ── Bold font helper ──
    def get_bold_font(font):
        if font == "Times-Roman":
            return "Times-Bold"
        elif font == "Helvetica":
            return "Helvetica-Bold"
        elif font == "Courier":
            return "Courier-Bold"
        else:
            return "Helvetica-Bold"

    font_name = style_cfg["font"]
    bold_font = get_bold_font(font_name)

    # ── Header alignment per template ──
    header_align_map = {
        "classic": TA_LEFT,
        "modern":  TA_CENTER,
        "compact": TA_LEFT,
        "minimal": TA_RIGHT
    }
    header_align = header_align_map.get(template_key, TA_LEFT)

    # ── Accent colors per template ──
    accent_map = {
        "classic": colors.HexColor("#1a1a1a"),
        "modern":  colors.HexColor("#2c5f9e"),
        "compact": colors.HexColor("#1a1a1a"),
        "minimal": colors.HexColor("#4a4a4a")
    }
    accent_color = accent_map.get(template_key, colors.black)

    # ── Heading border colors ──
    border_color_map = {
        "classic": colors.HexColor("#1a1a1a"),
        "modern":  colors.HexColor("#2c5f9e"),
        "compact": colors.HexColor("#555555"),
        "minimal": colors.HexColor("#aaaaaa")
    }
    border_color = border_color_map.get(template_key, colors.black)

    # ── Accent hex for inline markup ──
    accent_hex = {
        "classic": "#1a1a1a",
        "modern":  "#2c5f9e",
        "compact": "#1a1a1a",
        "minimal": "#4a4a4a"
    }.get(template_key, "#1a1a1a")

    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    story = []

    # ── STYLES ──
    name_style = ParagraphStyle(
        "Name",
        fontName=bold_font,
        fontSize=style_cfg.get("name_size", 16),
        textColor=accent_color,
        alignment=header_align,
        spaceAfter=4,
    )

    contact_style = ParagraphStyle(
        "Contact",
        fontName=font_name,
        fontSize=style_cfg["body_size"] - 0.5,
        textColor=colors.HexColor("#555555"),
        alignment=header_align,
        spaceBefore=2,
        spaceAfter=4,
    )

    heading_style = ParagraphStyle(
        "Heading",
        fontName=bold_font,
        fontSize=style_cfg["body_size"] + 1,
        textColor=accent_color,
        spaceBefore=12,    
        spaceAfter=2, 
    )

    body_style = ParagraphStyle(
        "Body",
        fontName=font_name,
        fontSize=style_cfg["body_size"],
        spaceAfter=6
    )

    bold_style = ParagraphStyle(
        "Bold",
        fontName=bold_font,
        fontSize=style_cfg["body_size"]
    )

    bullet_style = ParagraphStyle(
        "Bullet",
        fontName=font_name,
        fontSize=style_cfg["body_size"],
        leftIndent=style_cfg["bullet_indent"],
        firstLineIndent=0,
        spaceAfter=3
    )

    project_title_style = ParagraphStyle(
        "ProjectTitle",
        fontName=bold_font,
        fontSize=style_cfg["body_size"]
    )

    lines = text.split("\n")
    current_section = None
    section_elements = []
    bullets = []

    def flush_bullets_into_section():
        nonlocal bullets, section_elements
        if not bullets:
            return
        clean_items = []
        for b in bullets:
            text_content = re.sub(r"^[\s]*([0-9A-Za-z]+[\.\)]|[•\-\s]+)", "", b).strip()
            clean_items.append(
                ListItem(
                    Paragraph(text_content, bullet_style),
                    bulletText="•",
                    leftIndent=style_cfg["bullet_indent"],
                    bulletColor=accent_color,
                )
            )
        section_elements.append(
            ListFlowable(
                clean_items,
                bulletType='bullet',
                leftIndent=0,
                bulletIndent=0,
                spaceAfter=6
            )
        )
        bullets = []

    def flush_section():
        nonlocal section_elements
        if section_elements:
            story.append(KeepTogether(section_elements))
            section_elements = []

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # ── NAME ─_
        if i == 0:
            parts = line.split("|")
            name = parts[0].strip()
            contact_inline = " | ".join(p.strip() for p in parts[1:]) if len(parts) > 1 else ""

            story.append(Paragraph(name.upper(), name_style))
            if contact_inline:
                story.append(Paragraph(contact_inline, contact_style))
                story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
                story.append(Spacer(1, 4))
            i += 1
            continue

        # ── CONTACT LINE ──
        if i == 1 and not line.startswith("**"):
            contact_lines = [line]
            if i + 1 < len(lines) and lines[i+1].strip() and not lines[i+1].strip().startswith("**"):
                contact_lines.append(lines[i+1].strip())
                i += 1
            combined = " | ".join(contact_lines)
            story.append(Paragraph(combined, contact_style))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 4))
            i += 1
            continue

        # ── SECTION HEADING ──
        if line.startswith("**") and line.endswith("**"):
            flush_bullets_into_section()
            flush_section()
            current_section = line.replace("*", "")
            section_elements = [
                Paragraph(current_section, heading_style),
                HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc")),
                Spacer(1, 4)
            ]
            i += 1
            continue

        # ── EDUCATION ──
        if current_section == "EDUCATION":
            line1 = line
            line2 = lines[i + 1].strip() if i + 1 < len(lines) else ""
            section_elements.append(Paragraph(line1, bold_style))
            section_elements.append(Paragraph(line2, body_style))
            section_elements.append(Spacer(1, 8))
            i += 2
            continue

        # ── BULLETS ──
        bullet_match = re.match(r"^(\d+[\.\)]\s+|[A-Za-z][\.\)]\s+|[•\-]\s*)(.+)", line)
        if bullet_match:
            bullets.append(line)
            i += 1
            continue

        # ── PROJECT / WORK EXPERIENCE TITLE ──
        if current_section in ("PROJECTS", "WORK EXPERIENCE"):
            flush_bullets_into_section()
            section_elements.append(Paragraph(line, project_title_style))
            i += 1
            continue

        # ── TECHNICAL SKILLS ──
        if current_section == "TECHNICAL SKILLS" and ":" in line:
            left, right = line.split(":", 1)
            section_elements.append(
                Paragraph(
                    f'<font color="{accent_hex}"><b>{left}:</b></font> {right.strip()}',
                    body_style
                )
            )
            i += 1
            continue

        # ── NORMAL TEXT ──
        flush_bullets_into_section()
        section_elements.append(Paragraph(line, body_style))
        i += 1

    flush_bullets_into_section()
    flush_section()
    doc.build(story)


# 4. HTML RENDERING (FOR PREVIEW)
def render_resume_html(text, template_key):
    style = RESUME_TEMPLATES.get(template_key, RESUME_TEMPLATES["classic"])

    if not text or not isinstance(text, str):
        return "<p>No resume content available.</p>"

    # ── Per-template header alignment ──
    header_align = {
        "classic": "left",
        "modern":  "center",
        "compact": "left",
        "minimal": "right"
    }.get(template_key, "left")

    # ── Per-template accent colors ──
    accent_color = {
        "classic": "#1a1a1a",
        "modern":  "#2c5f9e",
        "compact": "#1a1a1a",
        "minimal": "#4a4a4a"
    }.get(template_key, "#1a1a1a")

    # ── Per-template heading border ──
    heading_border = {
        "classic": "border-bottom: 2px solid #1a1a1a;",
        "modern":  "border-bottom: 2px solid #2c5f9e;",
        "compact": "border-bottom: 1px solid #555;",
        "minimal": "border-bottom: 1px solid #aaa;"
    }.get(template_key, "border-bottom: 1px solid #333;")

    font          = style["font"].replace("-", " ")
    body_size     = style["body_size"]
    heading_size  = style["heading_size"]
    name_size     = style.get("name_size", 16)
    bullet_indent = style["bullet_indent"]

    lines = text.split("\n")
    html  = ""
    current_section = ""
    bullet_buffer   = []
    i = 0

    def flush_bullets():
        nonlocal html, bullet_buffer
        if bullet_buffer:
            html += (
                f'<ul style="margin: 4px 0 4px {bullet_indent}px;'
                f'padding: 0; list-style-type: disc;">'
            )
            for b in bullet_buffer:
                html += (
                    f'<li style="font-family:{font}; font-size:{body_size}px;'
                    f'margin-bottom:3px; padding-left:4px; color:#222;">{b}</li>'
                )
            html += '</ul>'
            bullet_buffer = []

    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # ── NAME ──
        if i == 0:
            parts = line.split("|")
            name = parts[0].strip()
            contact_inline = " | ".join(p.strip() for p in parts[1:]) if len(parts) > 1 else ""

            html += (
                f'<div style="font-size:{name_size}px; font-weight:bold;'
                f'font-family:{font}; text-align:{header_align};'
                f'text-transform:uppercase; color:{accent_color};'
                f'letter-spacing:1.5px; margin-bottom:4px;">'
                f'{name}</div>'
            )
            if contact_inline:
                html += (
                    f'<div style="font-size:{body_size - 0.5}px; color:#555;'
                    f'font-family:{font}; text-align:{header_align};'
                    f'margin-bottom:12px; line-height:1.6;">{contact_inline}</div>'
                    f'<hr style="border:none; border-top:1px solid #ccc; margin-bottom:10px;">'
                )
            i += 1
            continue

        # ── CONTACT LINE  ──
        if i == 1 and not line.startswith("**"):
            contact_lines = [line]
            if i + 1 < len(lines) and lines[i+1].strip() and not lines[i+1].strip().startswith("**"):
                contact_lines.append(lines[i+1].strip())
                i += 1
            combined = " | ".join(contact_lines)
            html += (
                f'<div style="font-size:{body_size - 0.5}px; color:#555;'
                f'font-family:{font}; text-align:{header_align};'
                f'margin-top:4px; margin-bottom:14px; line-height:1.6;">{combined}</div>'
                f'<hr style="border:none; border-top:1px solid #ccc; margin-bottom:12px;">'
            )
            i += 1
            continue

        # ── SECTION HEADING ──
        if line.startswith("**") and line.endswith("**"):
            flush_bullets()
            current_section = line.replace("*", "")
            html += (
                f'<div style="font-size:{heading_size}px; font-weight:bold;'
                f'font-family:{font}; color:{accent_color};'
                f'text-transform:uppercase; letter-spacing:1px;'
                f'margin-top:{style["section_gap"]}px; margin-bottom:2px;'
                f'padding-bottom:3px; {heading_border}">'
                f'{current_section}</div>'
                f'<div style="margin-bottom:6px;"></div>'
            )
            i += 1
            continue

        # ── EDUCATION ──
        if current_section == "EDUCATION":
            line1 = line
            line2 = lines[i + 1].strip() if i + 1 < len(lines) else ""
            html += (
                f'<div style="margin-bottom:8px;">'
                f'<div style="font-weight:bold; font-family:{font};'
                f'font-size:{body_size}px; color:#1a1a1a;">{line1}</div>'
                f'<div style="font-family:{font}; font-size:{body_size - 0.5}px;'
                f'color:#444;">{line2}</div>'
                f'</div>'
            )
            i += 2
            continue

        # ── BULLETS ──
        if line.startswith("•"):
            bullet_buffer.append(line[1:].strip())
            i += 1
            continue

        # ── PROJECT / WORK EXPERIENCE TITLE ──
        if current_section in ("PROJECTS", "WORK EXPERIENCE") and not line.startswith("•"):
            flush_bullets()
            html += (
                f'<div style="font-weight:bold; font-family:{font};'
                f'font-size:{body_size}px; color:#1a1a1a; margin-top:6px;">{line}</div>'
            )
            i += 1
            continue

        # ── TECHNICAL SKILLS ──
        if current_section == "TECHNICAL SKILLS" and ":" in line:
            left, right = line.split(":", 1)
            html += (
                f'<p style="font-family:{font}; font-size:{body_size}px; margin:3px 0;">'
                f'<strong style="color:{accent_color};">{left}:</strong>'
                f' {right.strip()}</p>'
            )
            i += 1
            continue

        # ── NORMAL TEXT ──
        flush_bullets()
        html += (
            f'<p style="font-family:{font}; font-size:{body_size}px;'
            f'margin:3px 0; color:#222;">{line}</p>'
        )
        i += 1

    flush_bullets()
    return html


def get_cover_letter(optimized_resume, job_description):
    url = "https://api.groq.com/openai/v1/chat/completions"

    prompt = f"""Write a professional cover letter based on the resume and job description below.

RULES:
- Start with candidate's name and contact from resume
- 3 paragraphs: Introduction, Why I am fit, Closing
- Tailor strongly to the job description
- Professional tone, no generic phrases
- Do NOT use placeholders like [Company Name] — extract company name from JD if available
- End with candidate's name

Resume:
{optimized_resume}

Job Description:
{job_description}"""

    headers = {
        "Authorization": f"Bearer {settings.GROQ_API_KEY}",
        "Content-Type": "application/json"
    }

    session = requests.Session()
    payload = {
        "model": "llama-3.3-70b-versatile",
        "messages": [
            {"role": "system", "content": "You are a professional cover letter writer."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.4,
        "max_tokens": 1024
    }

    response = session.post(url, json=payload, headers=headers, timeout=(15, 120))
    data = response.json()

    if "choices" not in data:
        raise Exception(data.get("error", "Groq API error"))

    return data["choices"][0]["message"]["content"]


def generate_cover_letter_pdf(text, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        leftMargin=60, rightMargin=60,
        topMargin=60, bottomMargin=60
    )

    styles_normal = ParagraphStyle(
        "Normal",
        fontName="Helvetica",
        fontSize=11,
        leading=18,
        spaceAfter=12
    )

    story = []
    for para in text.split("\n\n"):
        if para.strip():
            story.append(Paragraph(para.strip().replace("\n", "<br/>"), styles_normal))
            story.append(Spacer(1, 6))

    doc.build(story)